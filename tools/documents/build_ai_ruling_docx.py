from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs" / "reviews" / "AI对规则疑点的初步裁定.docx"
OUT.parent.mkdir(parents=True, exist_ok=True)

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(0.82)
sec.bottom_margin = Inches(0.78)
sec.left_margin = Inches(0.88)
sec.right_margin = Inches(0.88)
sec.header_distance = Inches(0.35)
sec.footer_distance = Inches(0.35)

NAVY = RGBColor(31, 55, 76)
BLUE = RGBColor(42, 92, 122)
GRAY = RGBColor(95, 104, 112)
LIGHT = "EAF1F5"
GOLD = RGBColor(170, 116, 38)

def set_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    if size is not None: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if color is not None: run.font.color.rgb = color
    if italic is not None: run.italic = italic

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.22

for name, size, color, before, after in [
    ("Heading 1", 16, NAVY, 14, 7),
    ("Heading 2", 12.5, BLUE, 9, 4),
    ("Heading 3", 11, NAVY, 7, 3),
]:
    st = styles[name]
    st.font.name = "Microsoft YaHei"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    st.font.size = Pt(size); st.font.bold = True; st.font.color.rgb = color
    st.paragraph_format.space_before = Pt(before); st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

for name in ("List Bullet", "List Number"):
    st = styles[name]
    st.font.name = "Microsoft YaHei"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    st.font.size = Pt(10.3)
    st.paragraph_format.space_after = Pt(3)
    st.paragraph_format.line_spacing = 1.18

# Running header/footer
h = sec.header.paragraphs[0]
h.text = "合掌网页版 · 规则裁定备忘录"
h.alignment = WD_ALIGN_PARAGRAPH.LEFT
set_font(h.runs[0], 8.5, True, GRAY)
f = sec.footer.paragraphs[0]
f.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = f.add_run("2026-08-01  ·  第 ")
set_font(r, 8.5, color=GRAY)
fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
f._p.append(fld)
r = f.add_run(" 页")
set_font(r, 8.5, color=GRAY)

# Masthead
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(3)
r = p.add_run("规则裁定备忘录")
set_font(r, 10, True, GOLD)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(5)
r = p.add_run("AI对规则疑点的初步裁定")
set_font(r, 24, True, NAVY)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(15)
r = p.add_run("对《合掌网页版游戏规则审计与待裁定清单（2026-08-01）》所列问题的逐项建议")
set_font(r, 11.5, color=GRAY)

for label, value in [
    ("审计对象", "《合掌网页版统一规则（2026-08-01规则稿）》"),
    ("裁定日期", "2026年8月1日"),
    ("文档性质", "AI初步裁定建议；供规则制定者确认，不自动替代正式规则"),
]:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(label + "："); set_font(r, 9.5, True, NAVY)
    r = p.add_run(value); set_font(r, 9.5, color=RGBColor(45,45,45))

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(11); p.paragraph_format.space_after = Pt(12)
pPr = p._p.get_or_add_pPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), LIGHT); pPr.append(shd)
r = p.add_run("裁定总原则　"); set_font(r, 10.5, True, NAVY)
r = p.add_run("同时出手的合法行动一经建立，不因随机结算顺序消失；严格区分“命中”“造成伤害”“本体失去生命”；技能可以改变通常规则，但不得产生无成本复制、无限递归或回溯试错。")
set_font(r, 10.5, color=RGBColor(35,45,52))

doc.add_heading("一、统一术语", level=1)
for t in [
    "承受一次攻击：攻击确认命中；占用多人每回合一次攻击名额。",
    "受到伤害：正数伤害实际落到该角色控制的秘偶、龙、护盾或本体。",
    "本体受到伤害：本体生命因伤害减少。",
    "失去生命：生命直接减少，不属于伤害，不能被护盾、防伤或减伤处理。",
    "死亡：生命不大于0后，在当前原子结算结束时确认。",
    "技能伤害不属于攻击，除非技能明确写成一次攻击。",
]: doc.add_paragraph(t, style="List Bullet")

sections = [
("二、多人判刀与帮助", [
("1. 攻击者候选及死亡后的攻击", ["乙进入甲的小刀候选范围。所有同时选择、合法支付且完成“攻击建立”的攻击均锁定；攻击者后来被命中乃至死亡，其攻击仍继续结算。", "死亡只阻止以后选择新行动，不追溯取消已经建立的行动。这样可避免随机排序变相决定谁有资格出刀。"]),
("2. 每回合一次承受限制", ["该限制只计算“攻击命中”，包括极光产生的多目标攻击；不包括月亮防反、猎人死亡伤害、太阳领域伤害、自损或其他纯技能伤害。", "已承受一次攻击者不再成为后续攻击的目标。", "真实伤害若来自攻击，仍占名额；若来自技能，则不占。", "技能伤害不占名额，也不能提前保护角色免受普通攻击。"]),
("3. 帮助选择次数", ["帮助资格按每把攻击独立产生，而非整个等级一次。角色可以对多把攻击分别选择帮助。", "帮助未使其成为最终承受者时，可以继续参与后续攻击；攻击一旦确认命中该角色，其名额被占用，不能再帮接、帮闪或帮吃；接／闪成功不是承受攻击，不占攻击名额。"]),
("4. 帮接、帮闪同时出现", ["接与闪分别成功；所有选择帮助且其行动确实能防住该攻击的角色，都取得个人成功记录。", "囚徒累计成功闪，龙骑士累计成功接。公共结果仍只有一次：该攻击被防住，并产生一次重开。"]),
("5. 全部候选者均使用接或闪", ["无需随机指定一人。所有具有有效防御的候选者分别取得相应的接／闪成功记录；若既有接又有闪，则两类均成功，但整把攻击只产生一次公共重开。"]),
("6. 帮助决定的信息时序", ["同一阶段必须秘密、同时决定：先同时决定帮接／帮闪并公布；若没有防御帮助成功，再同时决定帮吃；最后随机判刀。", "帮吃者可以知道帮接／帮闪阶段是否失败，但不能知道其他人的帮吃选择或随机结果。"]),
]),
("三、抵刀与攻击失效", [
("7. “同攻击”统一口径", ["统一为：至少两个敌对阵营各有一把最终等级相同、且最终等级不低于3的有效攻击。不要求武器同名。优先级表中的“同攻击抵刀”应改成“同最终等级攻击抵刀”。"]),
("8. 抵刀取消范围", ["满足条件时，该最终等级的全部攻击均被抵刀，包括第三阵营及同阵营额外打出的该等级攻击。", "这是等级级别的公共结果，而不是攻击两两配对；否则三队局会出现任意配对和剩余攻击选择问题。"]),
("9. 无法使用、失效、不能命中", ["无法使用：在攻击建立前排除，不支付，也不参与抵刀。", "打出后失效：已经支付并建立，但不能命中或造成伤害；仍参与抵刀。", "不能命中：仍是有效攻击并参与抵刀，只是在后续没有合法命中结果。", "因攻击建立后的技能而无效，费用不返还，除非技能明确写明返还。奇迹令敌方无法使用的判断应前移至攻击建立阶段；非法行动应退回重新选择合法行动。"]),
]),
("四、命中、伤害、护盾与技能双人", [
("10. “受到伤害”的默认含义", ["默认指正数伤害实际落到角色控制的秘偶、龙、护盾或本体；“本体受到伤害”才专指本体生命减少。", "通识者、审判者、窥秘人、梦魇等只写“受到伤害”时，秘偶、龙、护盾承伤也会触发；伤害被完全防止或减至0不触发。低血量死亡规则继续明确要求本体受到正数伤害。"]),
("11. 护盾", ["一次单段非真实正数伤害消耗一层护盾，并防止该段全部伤害，不论数值大小；0伤害不消耗护盾。", "明确标注的多段伤害逐段处理，每段至多消耗一层；每段只消耗最外一层。", "护盾承受并防止正数伤害，算角色受到伤害，但不算本体受伤或本体失去生命。“失去护盾”类技能在护盾消耗或被技能移除时触发。"]),
("12. 怪兽防伤后的后果", ["攻击命中事实保留：双人仍重开；多人仍占用承受攻击名额；锻造者仍不获得锻；触发“命中时”技能。", "不触发“受到／造成伤害时”技能，倒吊人不能发动。"]),
("13. 技能双人的死亡、计分与重置", ["采用小局制：一方死亡，对方得1分；双方同时死亡则各得1分。默认先到3分获胜，可在开局设置中调整。", "下一小局双方恢复角色印刷的初始生命、攻击、护盾及初始产生物，并清除锻、技能使用次数、永久减费、偷取／复制技能、标记、领域和跨回合状态。", "普通命中重开只清零锻并进入下一回合，不恢复生命或技能资源；死亡重置属于下一小局，范围明显更大。"]),
]),
("五、死亡与同时结算", [
("14. 死亡检查时机", ["采用“原子效果结束后死亡，已建立攻击不取消”。每把攻击的伤害及其直接响应结算完后检查死亡；死者立刻离场并触发死亡效果，但其在本次出手中已经建立、尚未轮到的攻击继续结算。", "死者不能再作新的主动选择；已建立攻击所附带的必要选择按建立时决定，无法预选时采用规则默认或随机合法选择。"]),
("15. 同时死亡技能顺序", ["先建立同一死亡批次，再依次处理：猎人死亡伤害、倒吊人回血、圣斗士杀人所得护盾、盗贼永久获得技能、占卜家替换秘偶、技能产生物清理。", "同类多个触发按固定座位顺序，从本轮起始玩家开始顺时针处理。新死亡进入下一死亡批次，不插入当前批次。盗贼取得的技能按该技能自身所属时机结算。"]),
("16. 技能失去生命", ["倒吊人、渔夫、先驱的自损不占承受攻击名额，不属于受到伤害，不触发伤害类技能，也不能被护盾、减伤、防伤或怪兽阻止，但可以导致死亡。"]),
]),
("六、高风险技能", [
("17. 魔术师", ["可指定任一仍存活、未退场的角色，包括原攻击者、目标、友方或敌方。替换攻击是技能建立的虚拟攻击，指定角色不支付锻，也不要求其有足够锻。", "继承原攻击的武器名、最终等级、目标、基础伤害、真实伤害、多目标及已附着的宣言／易伤状态；攻击力改用指定角色替换时的当前攻击力，只适用指定角色的伤害修正技能。", "不重新抵刀、不重新选择目标、不重新处理费用及“打出武器时”技能；重新处理最终攻击者与目标的怪兽、倒吊人等命中后响应。", "同一次原始攻击最多被魔术师替换一次；替换生成的攻击不能再次触发魔术师。"]),
("18. 黑夜回溯", ["本轮开始时建立完整快照，恢复生命、攻击力、锻、护盾、技能次数、永久减费、附加技能、死亡／退场状态、秘偶、龙、鱼钩、标记、领域、倒计时、赠予物，以及本轮内产生的得分和死亡连锁结果。", "不恢复本轮开始前的状态，也不能撤销已经产生的整轮胜负。玩家保留已知信息，但随机结果重新生成；回溯后从本轮第一回合开始重新选择。", "发动后立即失去黑夜技能；这一失去属于回溯外成本，不被快照恢复。"]),
("19. 盗贼的“剩余技能”", ["包括目标当前拥有的被动、主动和领域技能；不包括角色身份规则、初始生命／攻击、初始护盾、秘偶、龙等实体。", "有限次数以目标被盗时的剩余次数为上限，盗贼使用同步减少临时副本；目标恢复后保留被盗期间的消耗。", "不自动取得目标后来获得的新技能；永久获得以盗取时冻结的技能集合为准。角色初始化步骤不复制，除非其文字明确是可重复触发的被动。"]),
("20. 占卜家秘偶", ["秘偶生命归零后立即移除，对应技能同时消失。秘偶不是角色，不能被选择、标记、治疗或作为攻击目标。", "占卜家可正常使用对应被动、主动和领域；有限次数独立计数，从完整初始次数开始。", "替换时旧秘偶技能、次数、标记和由其持续维持的效果全部清除；新秘偶从完整初始技能状态开始，但不追溯已错过的“每轮开始前”时机。"]),
("21. 审判者追加对局", ["只有实际造成正数伤害才可发动；防止伤害不触发。追加对局中攻击命中、接成功或闪成功即结束；抵刀不结束。", "追加对局不触发角色技能，死亡暂不触发主场死亡技能；返回主场、同步生命后统一检查死亡。", "每个伤害事件，每名审判者最多发动一次并消耗一次次数。返回后从该伤害之后的造成伤害阶段继续。审判者因触发伤害处于0血时仍可发动一次追加对局。"]),
("22. 渔夫挑战", ["使用独立无技能双人小局，双方继承当前生命，锻从0开始。攻击命中的一方胜；接／闪者胜；攻击抵刀则继续。", "渔夫胜则移除鱼钩、杀死发动者且技能不生效；渔夫负则鱼钩保留、被挑战技能继续生效。", "同一主动技能只能被一个鱼钩挑战；多个渔夫按固定座位顺序取得声明权。领域不属于主动技能，不能被鱼钩挑战。"]),
("23. 奇迹宣言", ["三项效果作为一个宣言同时生效：自己的攻击不能命中；更低等级攻击不能命中自己；敌方同最终等级武器无法使用。", "宣言无独立次数限制，但必须随武器打出。时序为选择武器和宣言、检查合法性、对合法行动支付。被判无法使用者不支付并重新选择；奇迹自己的武器仍支付。"]),
("24. 极光多目标攻击", ["命中原判刀范围内每个敌对阵营各一名角色，而非全部角色。每队目标从该队合法候选者中随机产生，攻击者不能指定。", "各阵营分别处理帮助、承受者、魔术师、怪兽和伤害；某队接住／闪掉只保护该队分支。每个实际命中的角色分别占用名额。", "已承受攻击者不进入候选范围；某队无人可承受则跳过。整把多目标攻击全部分支结束后，最多产生一次公共重开。"]),
("25. 月亮防反", ["对本次原判刀范围内、其武器原本可能命中月亮的所有敌方武器使用者反击；已抵刀、失效、不能命中月亮或只攻击其他范围者不反击。", "“其攻击力”指各被反击者自己的当前攻击力。防反是月亮来源的技能伤害，不是攻击命中，不占攻击名额，也不触发怪兽或重开。"]),
("26. 先驱", ["被动为强制效果：每次即将受到正数伤害时，防止该伤害并失去1生命。主动已防止本回合伤害时，不再进入被动。", "多次伤害分别触发被动；失去生命不能被护盾或减伤处理。先驱自损死亡时，原攻击仍保持命中、占用名额并导致重开，但不造成伤害。"]),
("27. 领域、奇点与渔夫先后", ["顺序为：同时秘密声明领域并消耗次数；渔夫只拦截主动而不拦截领域；检查敌对领域对冲；未被对冲的奇点领域建立；奇点令敌方被动和主动失效，但不令领域失效。", "被对冲、封锁或挑战失败的已声明技能不返还次数。"]),
("28. 赠予敌人的效果", ["黑皇帝赠出的护盾仍属于黑皇帝技能产生物并记录来源；黑皇帝死亡时消失，失去护盾可触发圣斗士。", "牧羊人可以增加敌人的有限主动技能次数，因为原文没有阵营限制。治疗敌人是施术者造成的技能效果，但不是友方效果；治疗不属于伤害，也不撤销已发生的死亡。"]),
]),
("七、基础规则边界", [
("29. 1.5级攻击", ["普通防可防风暴之刃；接不能接风暴之刃。风暴之刃高于小刀、低于武士刀。", "风暴之刃与剪刀相遇时优先适用剪刀特殊克制；剪刀克制武士刀，因此也克制由武士刀改成的风暴之刃。多人抵刀只适用于最终等级不低于3，故1.5级不能抵刀。"]),
("30. 锻数上限", ["任何获得、返还后的可用锻都不能超过6。白塔返还是撤销已支付花费，不是新获得锻；在6锻支付后返还可恢复到6。", "免除花费则从未扣除，事件记录不同。守卫、观者等额外获得发生在已达6时，溢出直接消失。所有扣费先执行，再处理返还和获得；返还不得超过支付前锻数。"]),
("31. 计时结束", ["计时器属于网页版行动提交机制，不属于线下核心胜负规则。真人超时默认锻造；不能锻造者选择预设的最低风险合法行动，建议为防，仍非法则从合法行动中等概率随机。", "AI适用同一截止与默认规则。不应随机替玩家消耗锻，也不应以扣血、死亡惩罚。计时长度、动画延迟和AI思考时间应写入实现规范，而非核心规则正文。"]),
]),
]

for section_title, items in sections:
    doc.add_heading(section_title, level=1)
    for title, paras in items:
        doc.add_heading(title, level=2)
        for idx, text in enumerate(paras):
            style = "List Bullet" if len(paras) > 1 else None
            doc.add_paragraph(text, style=style)

doc.add_heading("八、统一规则稿的优先修订事项", level=1)
for t in [
    "将抵刀统一成“同最终等级攻击”，并把“无法使用”前移到攻击建立前。",
    "把一次承受上限明确限定为“攻击命中”，排除技能伤害。",
    "将死亡结算改成“每个原子效果后检查，但已建立攻击继续”。",
    "重写“受到伤害／本体受到伤害／失去生命”的术语。",
    "为魔术师增加“一次原始攻击最多替换一次”。",
    "为黑夜规定快照范围及不可回溯的发动成本。",
    "将极光明确为“每个敌对阵营一名目标、各分支独立防御、整把攻击一次重开”。",
    "明确奇点只封锁被动和主动，不封锁领域。",
]: doc.add_paragraph(t, style="List Bullet")

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(10)
r = p.add_run("结论：")
set_font(r, 10.5, True, NAVY)
r = p.add_run("本套裁定使多人局的随机性只决定“谁承受攻击”，不再决定“谁的行动还能存在”；同时将技能稳定绑定到命中、伤害或失去生命三个不同事件上。")
set_font(r, 10.5)

doc.add_heading("九、查阅材料", level=1)
for t in [
    "《合掌网页版统一规则（2026-08-01规则稿）》",
    "《合掌网页版游戏规则审计与待裁定清单（2026-08-01）》",
    "《合掌技能版（第二版）》",
    "《合掌学基础》",
    "《合掌游戏—速成攻略—终稿》",
]: doc.add_paragraph(t, style="List Bullet")

# keep headings with following content and avoid widows where possible
for p in doc.paragraphs:
    p.paragraph_format.widow_control = True

doc.core_properties.title = "AI对规则疑点的初步裁定"
doc.core_properties.subject = "合掌网页版规则审计初步裁定"
doc.core_properties.author = "OpenAI Codex"
doc.save(OUT)
print(OUT)
