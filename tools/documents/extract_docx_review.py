from pathlib import Path
from docx import Document
import zipfile

root = Path(__file__).resolve().parents[2]
excluded = "合掌技能版（第一版）_1784612659965.docx"
out = root / ".cache" / "docx-review"
out.mkdir(parents=True, exist_ok=True)

inputs = sorted((root / "docs" / "rules").rglob("*.docx")) + sorted((root / "docs" / "reviews").glob("*.docx"))
for path in inputs:
    if path.name == excluded:
        continue
    doc = Document(path)
    lines = [f"# FILE: {path.name}", "", "## BODY"]
    for i, p in enumerate(doc.paragraphs, 1):
        text = p.text.strip()
        if text:
            style_name = p.style.name if p.style is not None else "(none)"
            lines.append(f"P{i} [{style_name}]: {text}")
    for ti, table in enumerate(doc.tables, 1):
        lines.extend(["", f"## TABLE {ti}"])
        for ri, row in enumerate(table.rows, 1):
            cells = [c.text.replace("\n", " / ").strip() for c in row.cells]
            lines.append(f"R{ri}: " + " | ".join(cells))
    for si, section in enumerate(doc.sections, 1):
        h = " / ".join(p.text.strip() for p in section.header.paragraphs if p.text.strip())
        f = " / ".join(p.text.strip() for p in section.footer.paragraphs if p.text.strip())
        if h or f:
            lines.extend(["", f"## SECTION {si} HEADER/FOOTER", f"HEADER: {h}", f"FOOTER: {f}"])
    with zipfile.ZipFile(path) as z:
        names = set(z.namelist())
        for part in ("word/comments.xml", "word/footnotes.xml", "word/endnotes.xml"):
            if part in names:
                lines.extend(["", f"## RAW {part}", z.read(part).decode("utf-8", errors="replace")])
    (out / f"{path.stem}.txt").write_text("\n".join(lines), encoding="utf-8")
    print(path.name, len(doc.paragraphs), len(doc.tables))
