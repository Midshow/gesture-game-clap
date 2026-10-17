"""Synchronize the 2026-08-03 unified-rule DOCX body from its Markdown source.

The Markdown file is the editable source of truth.  This script keeps the
existing DOCX styles and layout, changing only body blocks whose visible text
differs from the Markdown representation.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph


def markdown_blocks(path: Path) -> list[str]:
    blocks: list[str] = []
    for raw in path.read_text(encoding="utf-8-sig").splitlines():
        text = raw.strip()
        if not text or text.startswith("<!--"):
            continue
        if re.fullmatch(r"\|?\s*:?-+:?\s*(\|\s*:?-+:?\s*)+\|?", text):
            continue
        text = re.sub(r"^#{1,6}\s+", "", text)
        text = re.sub(r"^[-*+]\s+", "", text)
        text = re.sub(r"^\d+\.\s+", "", text)
        if text.startswith("|") and text.endswith("|"):
            text = text[1:-1].strip()
        text = re.sub(r"\s*\|\s*", " | ", text)
        text = text.replace("**", "").replace("`", "")
        blocks.append(text)
    return blocks


def document_blocks(document: Document):
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            paragraph = Paragraph(child, document)
            if paragraph.text.strip():
                yield ("paragraph", paragraph, paragraph.text.strip())
        elif child.tag.endswith("}tbl"):
            table = Table(child, document)
            for row in table.rows:
                text = " | ".join(cell.text.strip().replace("\n", " / ") for cell in row.cells)
                yield ("row", row, text)


def normalized(text: str) -> str:
    return re.sub(r"\s+", "", text).replace("／", "/").replace("—", "-").replace("–", "-")


def replace_paragraph(paragraph: Paragraph, text: str) -> None:
    paragraph.text = text


def replace_row(row, text: str) -> None:
    values = [part.strip() for part in text.split("|")]
    if len(values) != len(row.cells):
        raise ValueError(f"Table column mismatch: expected {len(row.cells)}, got {len(values)}")
    for cell, value in zip(row.cells, values):
        cell.paragraphs[0].text = value
        for extra in list(cell.paragraphs[1:]):
            extra._element.getparent().remove(extra._element)


def main() -> None:
    root = Path(__file__).resolve().parent.parent
    rules = root / "docs" / "rules" / "unified"
    markdown = rules / "合掌网页版统一规则（2026-08-03）.md"
    docx = rules / "合掌网页版统一规则（2026-08-03）.docx"
    document = Document(docx)
    source = markdown_blocks(markdown)
    targets = list(document_blocks(document))
    if len(source) != len(targets):
        raise ValueError(f"Block count mismatch: Markdown {len(source)}, DOCX {len(targets)}")

    changed = 0
    for expected, (kind, target, current) in zip(source, targets):
        if normalized(expected) == normalized(current):
            continue
        if kind == "paragraph":
            replace_paragraph(target, expected)
        else:
            replace_row(target, expected)
        changed += 1

    document.save(docx)
    print(f"Synchronized {changed} changed blocks: {docx}")


if __name__ == "__main__":
    main()
